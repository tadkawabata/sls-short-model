"""
Recalculate all formulas in an .xlsx file via headless LibreOffice and report
any Excel error values. Adapted for a plain Ubuntu GitHub Actions runner
(no sandbox socket restrictions, so no AF_UNIX shim is needed).
"""
import json
import os
import subprocess
import sys
import tempfile
import time
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula

MACRO_FILENAME = "Module1.xba"
EXCEL_ERRORS = ["#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A"]
MAX_LOCATIONS = 100

RECALCULATE_MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>"""


def _existing_cells(ws):
    return [ws._cells[key] for key in sorted(ws._cells)]


def _stamp(path):
    st = os.stat(path)
    return st.st_mtime_ns, st.st_size


def recalc(filename, timeout=90):
    abs_path = str(Path(filename).absolute())
    if not Path(abs_path).exists():
        return {"error": f"File {filename} does not exist"}

    env = os.environ.copy()
    env["SAL_USE_VCLPLUGIN"] = "svp"

    with tempfile.TemporaryDirectory(prefix="lo_profile_") as profile_dir:
        profile_url = Path(profile_dir).as_uri()

        setup = subprocess.run(
            ["soffice", "--headless", "--terminate_after_init",
             f"-env:UserInstallation={profile_url}"],
            capture_output=True, timeout=timeout, env=env,
        )
        if setup.returncode != 0:
            return {"error": f"LibreOffice profile setup failed: {setup.stderr.decode(errors='replace')}"}

        macro_dir = Path(profile_dir) / "user" / "basic" / "Standard"
        macro_dir.mkdir(parents=True, exist_ok=True)
        (macro_dir / MACRO_FILENAME).write_text(RECALCULATE_MACRO)

        before = _stamp(abs_path)
        cmd = [
            "soffice", "--headless", "--norestore",
            f"-env:UserInstallation={profile_url}",
            "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application",
            abs_path,
        ]
        try:
            result = subprocess.run(cmd, capture_output=True, text=True, env=env, timeout=timeout)
        except subprocess.TimeoutExpired:
            return {"error": f"LibreOffice timed out after {timeout}s"}

        if result.returncode != 0:
            return {"error": f"LibreOffice failed: {result.stderr.strip()}"}

        if _stamp(abs_path) == before:
            return {"error": "LibreOffice exited cleanly but never rewrote the file"}

    wb = load_workbook(filename, data_only=True)
    error_details = {e: [] for e in EXCEL_ERRORS}
    total_errors = 0
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for cell in _existing_cells(ws):
            if isinstance(cell.value, str):
                for err in EXCEL_ERRORS:
                    if err in cell.value:
                        error_details[err].append(f"{sheet_name}!{cell.coordinate}")
                        total_errors += 1
                        break
    wb.close()

    wb_f = load_workbook(filename, data_only=False)
    formula_count = 0
    for sheet_name in wb_f.sheetnames:
        ws = wb_f[sheet_name]
        for cell in _existing_cells(ws):
            v = cell.value
            if isinstance(v, ArrayFormula):
                v = v.text
            if isinstance(v, str) and v.startswith("="):
                formula_count += 1
    wb_f.close()

    error_summary = {}
    for err_type, locations in error_details.items():
        if locations:
            entry = {"count": len(locations), "locations": locations[:MAX_LOCATIONS]}
            if len(locations) > MAX_LOCATIONS:
                entry["locations_truncated"] = len(locations) - MAX_LOCATIONS
            error_summary[err_type] = entry

    return {
        "status": "success" if total_errors == 0 else "errors_found",
        "total_errors": total_errors,
        "total_formulas": formula_count,
        "error_summary": error_summary,
    }


if __name__ == "__main__":
    filename = sys.argv[1]
    timeout = int(sys.argv[2]) if len(sys.argv) > 2 else 90
    result = recalc(filename, timeout)
    print(json.dumps(result, indent=2))
    sys.exit(1 if ("error" in result or result.get("status") == "errors_found") else 0)
